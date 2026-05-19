"""
Orquestador: lee Excel de placas desde S3, encola mensajes en SQS.

Columnas esperadas: numero_placas, cod_id, id (formato RUNT / Sufi).
"""

from __future__ import annotations

import argparse
import io
import json
import os
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import boto3
from openpyxl import load_workbook

from worker.src.plate_validation import is_valid_plate

ENDPOINT = os.getenv("AWS_ENDPOINT_URL", "http://localhost:4566")
REGION = os.getenv("AWS_DEFAULT_REGION", "us-east-1")
BUCKET_INPUT = os.getenv("S3_BUCKET_INPUT", "placas-proyecto-input")
QUEUE_NAME = os.getenv("SQS_QUEUE_NAME", "placas-jobs")
TABLE_JOBS = os.getenv("DYNAMODB_TABLE_JOBS", "placas-jobs")
BATCH_SIZE = int(os.getenv("SQS_BATCH_SIZE", "10"))


def _client(service: str):
    return boto3.client(
        service,
        endpoint_url=ENDPOINT,
        region_name=REGION,
        aws_access_key_id=os.getenv("AWS_ACCESS_KEY_ID", "test"),
        aws_secret_access_key=os.getenv("AWS_SECRET_ACCESS_KEY", "test"),
    )


def _header_map(ws) -> dict[str, int]:
    headers = [
        str(c.value).strip().lower() if c.value is not None else ""
        for c in next(ws.iter_rows(max_row=1))
    ]
    return {name: idx for idx, name in enumerate(headers)}


def read_rows_from_excel(data: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        raise ValueError("Excel sin hoja activa")

    headers = _header_map(ws)
    placa_col = headers.get("numero_placas", headers.get("placa", 0))
    cod_col = headers.get("cod_id", 1)
    id_col = headers.get("id", 2)

    rows: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or placa_col >= len(row):
            continue
        placa_val = row[placa_col]
        if placa_val is None:
            continue
        placa = str(placa_val).strip().upper()
        if not placa:
            continue
        cod_id = int(row[cod_col]) if cod_col < len(row) and row[cod_col] is not None else 1
        doc_id = row[id_col] if id_col < len(row) else ""
        if not is_valid_plate(placa):
            continue
        rows.append(
            {
                "placa": placa,
                "cod_id": cod_id,
                "documento_id": doc_id,
            }
        )
    wb.close()
    return rows


def enqueue_rows(job_id: str, rows: list[dict]) -> dict:
    sqs = _client("sqs")
    dynamodb = _client("dynamodb")
    queue_url = sqs.get_queue_url(QueueName=QUEUE_NAME)["QueueUrl"]
    sent = 0

    for i in range(0, len(rows), BATCH_SIZE):
        batch = rows[i : i + BATCH_SIZE]
        entries = [
            {
                "Id": str(j),
                "MessageBody": json.dumps({"job_id": job_id, **row}),
            }
            for j, row in enumerate(batch)
        ]
        sqs.send_message_batch(QueueUrl=queue_url, Entries=entries)
        sent += len(batch)

    dynamodb.put_item(
        TableName=TABLE_JOBS,
        Item={
            "job_id": {"S": job_id},
            "status": {"S": "QUEUED"},
            "total_placas": {"N": str(len(rows))},
            "queued_at": {"S": datetime.now(timezone.utc).isoformat()},
        },
    )
    return {"job_id": job_id, "total_placas": len(rows), "messages_sent": sent}


def run_job(job_id: str, s3_key: str) -> dict:
    s3 = _client("s3")
    data = s3.get_object(Bucket=BUCKET_INPUT, Key=s3_key)["Body"].read()
    rows = read_rows_from_excel(data)
    if not rows:
        raise ValueError(f"No se encontraron placas validas en s3://{BUCKET_INPUT}/{s3_key}")
    return enqueue_rows(job_id, rows)


def handler(event, context=None):
    job_id = event.get("job_id") or str(uuid.uuid4())
    result = run_job(job_id, event["s3_key"])
    return {"statusCode": 200, "body": json.dumps(result)}


def main():
    parser = argparse.ArgumentParser(description="Particionar Excel RUNT y encolar placas")
    parser.add_argument("--job-id", default=None)
    parser.add_argument("--s3-key", required=True)
    args = parser.parse_args()

    job_id = args.job_id or f"job-{uuid.uuid4().hex[:12]}"
    result = run_job(job_id, args.s3_key)
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
